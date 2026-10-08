#!/usr/bin/env python3
"""Convert the official CoFID 2021 workbook into the app's gram-based snapshot.

Requires openpyxl. This developer tool never ships in or runs from the app.
Download: https://www.gov.uk/government/publications/composition-of-foods-integrated-dataset-cofid
"""
import argparse
from collections import Counter
import hashlib
import json
import math
from pathlib import Path
import re

from openpyxl import load_workbook

COFID_ID_BASE = 1_000_000_000
COFID_DUPLICATE_BASE = 2_000_000_000
GROUPS = {
    'A': 'Cereals and cereal products', 'B': 'Milk and milk products',
    'C': 'Eggs', 'D': 'Vegetables', 'F': 'Fruit', 'G': 'Nuts and seeds',
    'H': 'Herbs and spices', 'J': 'Fish and fish products',
    'M': 'Meat and meat products', 'O': 'Fats and oils', 'P': 'Beverages',
    'Q': 'Alcoholic beverages', 'S': 'Sugars, preserves and snacks',
    'W': 'Soups, sauces and miscellaneous foods',
}


def numeric(value, maximum):
    """Missing and trace values have no reliable numeric amount; do not invent zero."""
    if value is None or isinstance(value, bool):
        return None
    try:
        result = float(value)
    except (ValueError, TypeError):
        return None
    return result if math.isfinite(result) and 0 <= result <= maximum else None


def food_id(code, name, duplicate=False):
    match = re.fullmatch(r'(\d{2})-(\d{3,4})', code)
    if not match:
        raise ValueError(f'Unexpected CoFID code: {code!r}')
    if duplicate:
        # The official workbook assigns 13-669 to both aubergine and watercress.
        # Keep both foods with deterministic IDs, rather than choosing one arbitrarily.
        digest = hashlib.sha256(f'{code}\0{name}'.encode('utf-8')).digest()
        return COFID_DUPLICATE_BASE + int.from_bytes(digest[:4], 'big')
    return COFID_ID_BASE + int(match[1]) * 10_000 + int(match[2])


def build(workbook):
    source = load_workbook(workbook, read_only=True, data_only=True)
    try:
        rows = list(source['1.3 Proximates'].iter_rows(min_row=4, values_only=True))
        factors = {row[0]: row for row in source['1.2 Factors'].iter_rows(min_row=4, values_only=True)}
    finally:
        source.close()
    codes = Counter(row[0] for row in rows)
    result, excluded = [], Counter()
    for row in rows:
        code, name, group = str(row[0]).strip(), str(row[1]).strip(), str(row[3]).strip()
        energy = numeric(row[12], 1000)
        if energy is None:
            excluded['missing_or_trace_energy'] += 1
            continue
        divisor, portions, note = 1.0, [], None
        if group.startswith('Q'):
            # CoFID alcoholic beverages are per 100 ml. Normalise only when the
            # workbook provides a plausible density, never assume 1 ml = 1 g.
            density = numeric(factors[code][8], 1.5)
            if density is None or density < 0.7:
                excluded['missing_or_implausible_alcohol_density'] += 1
                continue
            divisor = density
            portions = [{'name': '100 ml', 'grams': round(100 * density, 6)}]
            note = 'The UK source reports this drink per 100 ml. Gram values use its published density; the 100 ml portion reproduces the source serving.'
        nutrients = {key: numeric(row[index], 100) for key, index in [('protein', 9), ('fat', 10), ('carbohydrate', 11)]}
        converted_energy = energy / divisor
        converted = {key: round(value / divisor, 6) if value is not None else None for key, value in nutrients.items()}
        if converted_energy > 1000 or any(value is not None and value > 100 for value in converted.values()):
            excluded['out_of_range_converted_nutrients'] += 1
            continue
        entry = dict(id=food_id(code, name, codes[code] > 1), name=name,
            category=GROUPS[group[0]], kcalPer100g=round(converted_energy, 6),
            portions=portions, source='cofid', sourceCode=code, **converted)
        if note:
            entry['sourceNote'] = note
        result.append(entry)
    result.sort(key=lambda entry: (entry['name'].casefold(), entry['id']))
    if len({entry['id'] for entry in result}) != len(result):
        raise ValueError('Generated food ID collision')
    return result, dict(excluded), len(rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--output', type=Path, default=Path(__file__).resolve().parents[1] / 'FitnessApp/Resources/foods-uk.json')
    args = parser.parse_args()
    result, excluded, rows = build(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, separators=(',', ':'), allow_nan=False) + '\n')
    print(json.dumps(dict(sourceRows=rows, included=len(result), excluded=excluded,
        bytes=args.output.stat().st_size, sha256=hashlib.sha256(args.workbook.read_bytes()).hexdigest())))


if __name__ == '__main__':
    main()
